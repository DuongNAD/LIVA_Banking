import sys
import re
import json
import pptx

sys.stdout.reconfigure(encoding='utf-8')

tech_words = [
    'Rust', 'SIMD', 'AVX', 'AHash', 'u64', 'flatbuffer', 'DPAPI', 'Win32',
    'ReadDirectoryChanges', 'HMAC', 'SHA256', 'WAL', 'zero-copy', 'AST',
    'embedding', 'tokenizer', 'pipeline', 'IPC', 'daemon', 'mutex', 'JSON-RPC',
    'WebSocket', 'deserializ', 'backend', 'frontend', 'microsecond', 'nanosecond',
    'Tác tử'
]

prs = pptx.Presentation('liva_innostart_2026.pptx')
print("=== SCANNING PPTX NOTES FOR TECH TERMS ===")
for idx, slide in enumerate(prs.slides, 1):
    if slide.has_notes_slide:
        txt = slide.notes_slide.notes_text_frame.text
        for tw in tech_words:
            matches = re.findall(r'.{0,30}\b' + re.escape(tw) + r'\b.{0,30}', txt, re.IGNORECASE)
            if matches:
                for m in matches:
                    print(f"Slide {idx} Note matches '{tw}': {m.strip()}")

print("\n=== SCANNING PPTX SLIDE CONTENT FOR TECH TERMS ===")
for idx, slide in enumerate(prs.slides, 1):
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                txt = p.text
                for tw in tech_words:
                    matches = re.findall(r'.{0,30}\b' + re.escape(tw) + r'\b.{0,30}', txt, re.IGNORECASE)
                    if matches:
                        for m in matches:
                            print(f"Slide {idx} Shape matches '{tw}': {m.strip()}")

print("\n=== SCANNING INDEX.HTML FOR TECH TERMS ===")
with open('index.html', 'r', encoding='utf-8') as f:
    html_content = f.read()

for tw in tech_words:
    matches = re.findall(r'.{0,30}\b' + re.escape(tw) + r'\b.{0,30}', html_content, re.IGNORECASE)
    # Filter out CSS/JS identifiers if needed, but show matches
    real_matches = [m for m in matches if not ('class=' in m)]
    if real_matches:
        print(f"HTML matches for '{tw}': count={len(real_matches)}")
        for m in real_matches[:5]:
            print(f"   snippet: {m.strip()}")
