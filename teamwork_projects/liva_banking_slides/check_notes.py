import sys
import pptx
import re
import json

sys.stdout.reconfigure(encoding='utf-8')

prs = pptx.Presentation('liva_innostart_2026.pptx')
with open('index.html', 'r', encoding='utf-8') as f:
    html = f.read()

notes_match = re.search(r'window\.SPEAKER_NOTES\s*=\s*(\{.+?\});\s*</script>', html, re.DOTALL)
html_notes = json.loads(notes_match.group(1))

for idx in range(1, 26):
    slide = prs.slides[idx - 1]
    pptx_note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    html_note_obj = html_notes.get(str(idx), {})
    spoken = html_note_obj.get("spoken", "")
    takeaway = html_note_obj.get("takeaway", "")
    punchline = html_note_obj.get("punchline", "")
    
    print(f"=== SLIDE {idx} ===")
    print(f"  PPTX Note Length: {len(pptx_note)} chars")
    print(f"  HTML Spoken Length: {len(spoken)} chars")
    print(f"  HTML Takeaway: {takeaway[:80]}...")
    print(f"  HTML Punchline: {punchline[:80]}...")
