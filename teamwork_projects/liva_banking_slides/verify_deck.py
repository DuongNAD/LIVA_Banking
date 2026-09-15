#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
verify_deck.py
Automated 27-Slide Deck Integrity, Structure, Speaker Notes, and Jargon Audit.
Audits index.html, liva_innostart_2026.pptx, and liva_innostart_2026.pdf:
- HTML slide count = 27, PPTX slide count = 27, PDF page count = 27.
- HTML speaker notes count = 27, PPTX speaker notes count = 27.
- 0 occurrences of banned engineer jargon across HTML, PPTX, and PDF text.
"""

import os
import sys
import re
import json
import pptx
import pypdf

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

def verify():
    script_dir = os.path.dirname(os.path.abspath(__file__))
    pptx_path = os.path.join(script_dir, 'liva_innostart_2026.pptx')
    pdf_path = os.path.join(script_dir, 'liva_innostart_2026.pdf')
    html_path = os.path.join(script_dir, 'index.html')

    print("============================================================")
    print("LIVA BANKING HARNESS — 27-SLIDE DECK INTEGRITY AUDIT")
    print("============================================================")

    errors = []

    # 1. Verify File Existence
    for p in [pptx_path, pdf_path, html_path]:
        if not os.path.exists(p):
            errors.append(f"Missing file: {p}")
        else:
            print(f"[*] Found artifact: {os.path.basename(p)} ({os.path.getsize(p):,} bytes)")

    if errors:
        for e in errors:
            print(f"[!] ERROR: {e}")
        sys.exit(1)

    # 2. Check Slide Counts
    prs = pptx.Presentation(pptx_path)
    pptx_slide_count = len(prs.slides)
    print(f"\n[*] PPTX Slide Count: {pptx_slide_count} / 27")
    if pptx_slide_count != 27:
        errors.append(f"PPTX slide count mismatch: {pptx_slide_count} != 27")

    with open(html_path, 'r', encoding='utf-8') as f:
        html_content = f.read()

    slide_matches = re.findall(r'<section\s+class="slide[^"]*"\s+id="slide-(\d+)"', html_content)
    html_slide_count = len(slide_matches)
    print(f"[*] HTML Slide Count: {html_slide_count} / 27")
    if html_slide_count != 27:
        errors.append(f"HTML slide count mismatch: {html_slide_count} != 27")

    pdf_reader = pypdf.PdfReader(pdf_path)
    pdf_page_count = len(pdf_reader.pages)
    print(f"[*] PDF Page Count: {pdf_page_count} / 27")
    if pdf_page_count != 27:
        errors.append(f"PDF page count mismatch: {pdf_page_count} != 27")

    # 3. Check Speaker Notes Counts
    # HTML Speaker Notes
    notes_match = re.search(r'window\.SPEAKER_NOTES\s*=\s*(\{.+?\});\s*</script>', html_content, re.DOTALL)
    html_notes_count = 0
    if notes_match:
        notes_data = json.loads(notes_match.group(1))
        html_notes_count = len(notes_data)
        print(f"[*] HTML window.SPEAKER_NOTES keys: {html_notes_count} / 27")
        if html_notes_count != 27:
            errors.append(f"HTML speaker notes count mismatch: {html_notes_count} != 27")
    else:
        errors.append("Could not parse window.SPEAKER_NOTES from index.html")

    # PPTX Speaker Notes
    pptx_notes_count = 0
    for idx, slide in enumerate(prs.slides, 1):
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            txt = slide.notes_slide.notes_text_frame.text.strip()
            if len(txt) > 50:
                pptx_notes_count += 1
    print(f"[*] PPTX Embedded Speaker Notes: {pptx_notes_count} / 27")
    if pptx_notes_count != 27:
        errors.append(f"PPTX speaker notes count mismatch: {pptx_notes_count} != 27")

    # 4. Check Banned Engineer Jargon Across HTML, PPTX, and PDF
    deep_jargon = [
        'SIMD', 'AVX', 'AHash', 'Rust', 'deserializ', 'u64', 'Disentangle', 
        'DPAPI', 'Win32', 'ReadDirectoryChanges', 'HMAC', 'SHA256', 'Flatbuffers',
        'zero-copy', 'AST'
    ]

    print("\n--- CHECKING BANNED JARGON IN PPTX ---")
    pptx_jargon_found = []
    for idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    txt = p.text
                    for term in deep_jargon:
                        if re.search(r'\b' + re.escape(term) + r'\b', txt, re.IGNORECASE):
                            pptx_jargon_found.append((idx, term, txt))

    if pptx_jargon_found:
        for s_idx, term, txt in pptx_jargon_found:
            print(f"  [VIOLATION] PPTX Slide {s_idx} contains '{term}': {txt}")
            errors.append(f"PPTX Slide {s_idx} jargon violation: {term}")
    else:
        print("[✓] No deep technical jargon found in PPTX shapes.")

    print("\n--- CHECKING BANNED JARGON IN HTML ---")
    html_jargon_found = []
    html_slides = re.findall(r'<section\s+class="slide[^"]*"\s+id="slide-(\d+)"[^>]*>(.*?)</section>', html_content, re.DOTALL)
    for s_num, s_html in html_slides:
        plain_text = re.sub(r'<[^>]+>', ' ', s_html)
        for term in deep_jargon:
            if re.search(r'\b' + re.escape(term) + r'\b', plain_text, re.IGNORECASE):
                html_jargon_found.append((s_num, term))

    if html_jargon_found:
        for s_num, term in html_jargon_found:
            print(f"  [VIOLATION] HTML Slide {s_num} contains '{term}'")
            errors.append(f"HTML Slide {s_num} jargon violation: {term}")
    else:
        print("[✓] No deep technical jargon found in HTML slide surfaces.")

    print("\n--- CHECKING BANNED JARGON IN PDF ---")
    pdf_jargon_found = []
    for idx, page in enumerate(pdf_reader.pages, 1):
        txt = page.extract_text()
        for term in deep_jargon:
            if re.search(r'\b' + re.escape(term) + r'\b', txt, re.IGNORECASE):
                pdf_jargon_found.append((idx, term))

    if pdf_jargon_found:
        for p_idx, term in pdf_jargon_found:
            print(f"  [VIOLATION] PDF Page {p_idx} contains '{term}'")
            errors.append(f"PDF Page {p_idx} jargon violation: {term}")
    else:
        print("[✓] No deep technical jargon found in PDF extracted text.")

    # 5. Check Slide 14 Specific Guardrail (Zero Stale Roadmap Bullets)
    print("\n--- CHECKING SLIDE 14 STALE ROADMAP BULLETS ---")
    s14_pptx_clean = True
    if len(prs.slides) >= 14:
        s14 = prs.slides[13]
        for shape in s14.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if 'Chạy trên máy tính nội bộ' in p.text:
                        print(f"[VIOLATION] PPTX Slide 14 contains misleading bullet: '{p.text}'")
                        errors.append("Slide 14 contains stale roadmap bullet")
                        s14_pptx_clean = False
    if s14_pptx_clean:
        print("[✓] PPTX Slide 14 verified clean of stale roadmap bullets.")

    # Final status
    print("\n============================================================")
    if errors:
        print(f"[✗] VERIFICATION FAILED WITH {len(errors)} ERROR(S):")
        for e in errors:
            print(f"    - {e}")
        print("============================================================\n")
        sys.exit(1)
    else:
        print("[✓] ALL AUDIT CHECKS PASSED WITH 100% EXCELLENCE!")
        print("    - HTML Slides: 27 | PPTX Slides: 27 | PDF Pages: 27")
        print("    - HTML Speaker Notes: 27 | PPTX Speaker Notes: 27")
        print("    - Banned Jargon in HTML: 0 | PPTX: 0 | PDF: 0")
        print("============================================================\n")
        return True

if __name__ == '__main__':
    verify()

